#!/usr/bin/env python3
"""Parse nixlbench NIXLBENCH_RESULT / NIXLBENCH_CPU lines from log files and
fill the RDMA columns (Q-AA) of "Performance Plan for S3.xlsx".

Usage:
    python3 fill_sheet.py <xlsx_path> <log_file> [<log_file> ...]

Results are matched to spreadsheet rows by (op, object_size_MiB, threads).
Latency us -> ms. CPU: prefers whole-node mpstat %usr/%sys (NIXLBENCH_CPU);
falls back to getrusage process CPU normalized by core count.
"""
import re
import sys
import openpyxl

# RDMA column group (row 21/22 headers): Q..AA
COL = {
    "mibps": "R",
    "objps": "S",
    "avg": "T",
    "p90": "U",
    "p99": "V",
    "slowest": "W",
    "fastest": "X",
    "stddev": "Y",
    "usr": "Z",
    "sys": "AA",
}
FIRST_ROW = 23   # sheet row 23 == test 1
LAST_ROW = 56    # sheet row 56 == test 34


def parse_kv(line):
    return dict(re.findall(r"(\w+)=([-\w.]+)", line))


def parse_logs(paths):
    """Return dict keyed by (op, block_size, threads, dev) -> parsed kv.

    A NIXLBENCH_CPU line (whole-node mpstat %usr/%sys emitted by run_bench.sh)
    is attached to the most recent NIXLBENCH_RESULT in the same cell block.
    """
    results = {}
    last_key = None
    for p in paths:
        with open(p) as f:
            for line in f:
                if "NIXLBENCH_RESULT" in line:
                    kv = parse_kv(line)
                    if "block_size" not in kv:
                        continue
                    key = (
                        kv.get("op"),
                        int(kv["block_size"]),
                        int(kv.get("threads", 0)),
                        int(kv.get("initiator_dev", 0)),
                    )
                    results[key] = kv  # last wins
                    last_key = key
                elif "NIXLBENCH_CPU" in line and last_key is not None:
                    cpu = parse_kv(line)
                    if "usr" in cpu:
                        results[last_key]["mpstat_usr"] = cpu["usr"]
                    if "sys" in cpu:
                        results[last_key]["mpstat_sys"] = cpu["sys"]
    return results


def build_row_index(ws):
    """Map (op, block_size, threads) -> (row, expected_clients)."""
    idx = {}
    for r in range(FIRST_ROW, LAST_ROW + 1):
        size_mib = ws[f"C{r}"].value
        threads = ws[f"D{r}"].value
        op_cell = ws[f"Q{r}"].value  # RDMA op column
        if size_mib is None or threads is None or op_cell is None:
            continue
        op = "WRITE" if str(op_cell).strip().upper() == "PUT" else "READ"
        block_size = int(size_mib) * 1024 * 1024
        idx[(op, block_size, int(threads))] = (r, ws[f"E{r}"].value)
    return idx


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    xlsx, logs = sys.argv[1], sys.argv[2:]
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Performance plan"]

    results = parse_logs(logs)
    row_idx = build_row_index(ws)

    filled = 0
    for (op, bs, threads, dev), kv in sorted(results.items()):
        target = row_idx.get((op, bs, threads))
        if target is None:
            print(f"WARN: no sheet row for op={op} size={bs>>20}MiB threads={threads}")
            continue
        r, expected_clients = target
        if expected_clients is not None and int(expected_clients) != dev:
            print(f"WARN: row {r} (test {r-22}) expects {expected_clients} clients "
                  f"but run used {dev}; filling anyway")

        def num(k):
            return float(kv[k]) if k in kv else None

        if "mpstat_usr" in kv:
            usr, sysp = num("mpstat_usr"), num("mpstat_sys")
        else:
            wall = num("wall_s") or 0.0
            ncpu = num("ncpu") or 1.0
            usr = (num("cpu_usr_s") / wall / ncpu * 100) if wall else None
            sysp = (num("cpu_sys_s") / wall / ncpu * 100) if wall else None

        vals = {
            "mibps": num("bw_mibps"),
            "objps": num("obj_per_s"),
            "avg": (num("lat_avg_us") or 0) / 1000.0 if "lat_avg_us" in kv else None,
            "p90": (num("lat_p90_us") or 0) / 1000.0 if "lat_p90_us" in kv else None,
            "p99": (num("lat_p99_us") or 0) / 1000.0 if "lat_p99_us" in kv else None,
            "slowest": (num("lat_max_us") or 0) / 1000.0 if "lat_max_us" in kv else None,
            "fastest": (num("lat_min_us") or 0) / 1000.0 if "lat_min_us" in kv else None,
            "stddev": (num("lat_stddev_us") or 0) / 1000.0 if "lat_stddev_us" in kv else None,
            "usr": usr,
            "sys": sysp,
        }
        for name, v in vals.items():
            if v is not None:
                ws[f"{COL[name]}{r}"] = round(v, 4)
        filled += 1
        print(f"row {r} (test {r-22}): op={op} size={bs>>20}MiB threads={threads} clients={dev}")

    wb.save(xlsx)
    print(f"\nFilled {filled} row(s). Saved {xlsx}")


if __name__ == "__main__":
    main()
